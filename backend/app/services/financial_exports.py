"""PDF and Excel exports for deterministic financial calculation results."""

from dataclasses import asdict, is_dataclass
from decimal import Decimal
from enum import Enum
from io import BytesIO
import os
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


def _plain(value: Any) -> Any:
    if is_dataclass(value):
        return asdict(value)
    if isinstance(value, Enum):
        return value.value
    return value


def _font_file() -> Path:
    windows_dir = Path(os.environ.get("WINDIR", r"C:\Windows"))
    candidates = (
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        windows_dir / "Fonts" / "arial.ttf",
    )
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise RuntimeError("A Unicode TrueType font is required to create PDF exports")


def _register_pdf_font() -> str:
    font_name = "FinAdvisorUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, str(_font_file())))
        pdfmetrics.registerFontFamily(
            font_name,
            normal=font_name,
            bold=font_name,
            italic=font_name,
            boldItalic=font_name,
        )
    return font_name


def build_pdf_report(
    results: dict[str, Any],
    *,
    business_plan: str | None = None,
    risks: list[str] | None = None,
    bank_readiness_score: int | None = None,
    bank_readiness_score_note: str | None = None,
    bank_readiness_components: list[dict[str, Any]] | None = None,
) -> bytes:
    """Render the provided calculation result as a Unicode PDF in memory."""
    font_name = _register_pdf_font()
    buffer = BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="FinAdvisor — moliyaviy reja hisoboti",
        author="FinAdvisor",
    )
    base_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FinAdvisorTitle",
        parent=base_styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=24,
        textColor=colors.HexColor("#064e3b"),
        alignment=TA_LEFT,
        spaceAfter=5 * mm,
    )
    heading_style = ParagraphStyle(
        "FinAdvisorHeading",
        parent=base_styles["Heading2"],
        fontName=font_name,
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#065f46"),
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
    )
    cell_style = ParagraphStyle(
        "FinAdvisorCell",
        parent=base_styles["BodyText"],
        fontName=font_name,
        fontSize=8,
        leading=11,
    )
    story: list[Any] = [
        Paragraph("FinAdvisor — moliyaviy reja hisoboti", title_style),
        Paragraph(
            "Hisobot moliya dvigateli qaytargan qiymatlardan tuzildi. "
            "Kiritilmagan ma’lumotlar taxmin qilinmadi.",
            cell_style,
        ),
        Spacer(1, 4 * mm),
    ]

    sections = (
        ("Boshlang‘ich xarajatlar", results.get("capex")),
        ("Oylik operatsion xarajatlar", results.get("monthly_opex")),
        ("Yillik soliqlar", results.get("annual_taxes")),
        ("Zararsizlik nuqtasi", results.get("break_even")),
        ("Investitsiya daromadliligi", results.get("roi")),
    )
    for title, section in sections:
        if section is None:
            continue
        section = _plain(section)
        story.append(Paragraph(escape(title), heading_style))
        rows = [["Ko‘rsatkich", "Qiymat"]]
        rows.extend(_pdf_flat_rows(section, cell_style))
        story.append(_pdf_table(rows))

    if business_plan:
        story.append(Paragraph("Biznes reja izohi", heading_style))
        for paragraph_text in business_plan.splitlines():
            if paragraph_text.strip():
                story.append(Paragraph(escape(paragraph_text), cell_style))
                story.append(Spacer(1, 2 * mm))
    if risks:
        story.append(Paragraph("E’tibor berish kerak bo‘lgan jihatlar", heading_style))
        for risk in risks:
            story.append(Paragraph(f"• {escape(risk)}", cell_style))
    if bank_readiness_score is not None:
        story.append(Paragraph("Reja ma’lumotlari to‘liqligi", heading_style))
        story.append(Paragraph(f"{bank_readiness_score} / 100", cell_style))
        if bank_readiness_score_note:
            story.append(Paragraph(escape(bank_readiness_score_note), cell_style))
        if bank_readiness_components:
            score_rows = [["Mezon", "Ball", "Holat"]]
            score_rows.extend([
                [
                    Paragraph(escape(str(item["label"])), cell_style),
                    f"{item['earned']} / {item['possible']}",
                    "To‘liq" if item["complete"] else "Yetishmaydi",
                ]
                for item in bank_readiness_components
            ])
            story.append(_pdf_table(score_rows))

    loan_schedule = _plain(results.get("loan_schedule"))
    if loan_schedule is not None:
        installments = loan_schedule.get("installments", [])
        story.append(PageBreak())
        story.append(Paragraph("Kredit to‘lovlari grafigi", heading_style))
        story.append(
            Paragraph(
                f"Usul: {escape(str(loan_schedule.get('method', '')))} · "
                f"Asosiy qarz: {escape(str(loan_schedule.get('principal', '')))} so‘m · "
                f"Yillik stavka: {escape(str(loan_schedule.get('annual_interest_rate', '')))}%",
                cell_style,
            )
        )
        rows = [["Oy", "Qoldiq boshida", "Asosiy qarz", "Foiz", "Jami", "Qoldiq oxirida"]]
        for row in installments:
            rows.append([
                str(row["period"]),
                str(row["opening_balance"]),
                str(row["principal_payment"]),
                str(row["interest_payment"]),
                str(row["total_payment"]),
                str(row["closing_balance"]),
            ])
        story.append(_pdf_table(rows, compact=True))
        story.append(Spacer(1, 3 * mm))
        totals = [
            ["Jami asosiy qarz", str(loan_schedule.get("total_principal", ""))],
            ["Jami foiz", str(loan_schedule.get("total_interest", ""))],
            ["Jami to‘lov", str(loan_schedule.get("total_payment", ""))],
        ]
        story.append(_pdf_table([["Kredit jami", "So‘m"], *totals]))

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor("#64748b"))
        canvas.drawString(18 * mm, 10 * mm, "FinAdvisor · Moliyaviy hisob-kitob")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Sahifa {doc.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _pdf_flat_rows(section: Any, style: ParagraphStyle, prefix: str = "") -> list[list[Any]]:
    rows: list[list[Any]] = []
    if isinstance(section, dict):
        for key, value in section.items():
            label = f"{prefix}{_label(key)}"
            if isinstance(value, (dict, list, tuple)):
                nested_rows = _pdf_flat_rows(value, style, f"{label} · ")
                rows.extend(nested_rows)
            else:
                rows.append([
                    Paragraph(escape(label), style),
                    Paragraph(escape(_display(value)), style),
                ])
    elif isinstance(section, (list, tuple)):
        for index, value in enumerate(section, 1):
            rows.extend(_pdf_flat_rows(value, style, f"{prefix}{index} · "))
    return rows


def _display(value: Any) -> str:
    if value is None:
        return "Mavjud emas"
    if isinstance(value, Enum):
        return str(value.value)
    return str(value)


def _label(key: str) -> str:
    labels = {
        "equipment_items": "Uskunalar ro‘yxati",
        "equipment_total": "Uskunalar jami",
        "initial_investment": "Boshlang‘ich investitsiya",
        "payroll": "Ish haqi fondi",
        "rent": "Ijara",
        "utilities": "Kommunal to‘lovlar",
        "total": "Jami",
        "regime": "Soliq rejimi",
        "revenue_excluding_vat": "Yillik tushum (QQSsiz)",
        "vat_rate_applied": "Qo‘llangan QQS stavkasi",
        "output_vat": "Hisoblangan QQS",
        "input_vat_credit": "Hisobga olinadigan kiruvchi QQS",
        "unused_input_vat_credit": "Ishlatilmagan kiruvchi QQS",
        "vat_payable": "To‘lanadigan QQS",
        "taxable_profit": "Soliqqa tortiladigan foyda",
        "profit_tax_rate_applied": "Foyda solig‘i stavkasi",
        "profit_tax": "Foyda solig‘i",
        "turnover_tax_rate_applied": "Aylanmadan soliq stavkasi",
        "turnover_tax": "Aylanmadan soliq",
        "annual_gross_payroll": "Yillik ish haqi fondi",
        "personal_income_tax_withheld": "Ushlab qolinadigan JShDS",
        "employer_social_tax": "Ish beruvchi ijtimoiy solig‘i",
        "break_even_revenue": "Zararsizlik nuqtasidagi tushum",
        "break_even_total_units": "Zararsizlik nuqtasidagi mahsulot soni",
        "contribution_margin_ratio": "Qoplash marjasi ulushi",
        "annual_net_income_after_tax": "Yillik sof daromad (soliqdan keyin)",
        "roi_percent": "ROI foizda",
        "principal": "Asosiy qarz",
        "annual_interest_rate": "Yillik foiz stavkasi",
        "total_principal": "Jami asosiy qarz",
        "total_interest": "Jami foiz",
        "total_payment": "Jami to‘lov",
        "period": "Oy",
        "opening_balance": "Qoldiq boshida",
        "principal_payment": "Asosiy qarz to‘lovi",
        "interest_payment": "Foiz to‘lovi",
        "closing_balance": "Qoldiq oxirida",
    }
    return labels.get(key, key.replace("_", " "))


def _pdf_table(rows: list[list[Any]], compact: bool = False) -> Table:
    if compact or len(rows[0]) == 6:
        widths = [24 * mm, 29 * mm, 25 * mm, 23 * mm, 23 * mm, 31 * mm]
    elif len(rows[0]) == 3:
        widths = [85 * mm, 35 * mm, 35 * mm]
    else:
        widths = [105 * mm, 75 * mm]
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#065f46")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, -1), "FinAdvisorUnicode"),
        ("FONTSIZE", (0, 0), (-1, -1), 7 if compact else 8),
        ("LEADING", (0, 0), (-1, -1), 10),
        ("BACKGROUND", (0, 1), (-1, -1), colors.HexColor("#f8fafc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#e2e8f0")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    return table


def build_excel_report(
    results: dict[str, Any],
    *,
    business_plan: str | None = None,
    risks: list[str] | None = None,
    bank_readiness_score: int | None = None,
    bank_readiness_score_note: str | None = None,
    bank_readiness_components: list[dict[str, Any]] | None = None,
) -> bytes:
    """Create an XLSX workbook with engine values and visible formula notes."""
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Moliyaviy natijalar"
    _style_sheet_title(summary, "FinAdvisor — moliyaviy reja")
    summary.append(["Bo‘lim", "Ko‘rsatkich", "Hisoblangan qiymat"])
    _style_header(summary[2])

    sections = (
        ("CAPEX", results.get("capex")),
        ("OPEX / oy", results.get("monthly_opex")),
        ("Soliqlar / yil", results.get("annual_taxes")),
        ("Zararsizlik", results.get("break_even")),
        ("ROI", results.get("roi")),
        ("Kredit", results.get("loan_schedule")),
    )
    for section_name, raw_section in sections:
        if raw_section is None:
            continue
        section = _plain(raw_section)
        for label, value in _flatten(section):
            if label == "installments":
                continue
            summary.append([section_name, label, _excel_value(value)])

    if business_plan or risks or bank_readiness_score is not None:
        plan_sheet = workbook.create_sheet("Biznes reja")
        _style_sheet_title(plan_sheet, "Biznes reja izohi")
        plan_sheet.append(["Qism", "Matn"])
        _style_header(plan_sheet[2])
        if business_plan:
            plan_sheet.append(["Reja", _excel_value(business_plan)])
        if bank_readiness_score is not None:
            plan_sheet.append(["Reja to‘liqligi", bank_readiness_score])
        if bank_readiness_score_note:
            plan_sheet.append(["Izoh", _excel_value(bank_readiness_score_note)])
        for item in bank_readiness_components or []:
            plan_sheet.append([
                _excel_value(item["label"]),
                f"{item['earned']} / {item['possible']}",
            ])
        for index, risk in enumerate(risks or [], 1):
            plan_sheet.append([f"Xavf {index}", _excel_value(risk)])

    loan_schedule = _plain(results.get("loan_schedule"))
    if loan_schedule is not None:
        loan_sheet = workbook.create_sheet("Kredit grafigi")
        _style_sheet_title(loan_sheet, "Kredit to‘lovlari · moliya dvigateli natijalari")
        loan_sheet.append([
            "Oy", "Qoldiq boshida", "Asosiy qarz", "Foiz", "Jami to‘lov", "Qoldiq oxirida",
        ])
        _style_header(loan_sheet[2])
        for row in loan_schedule.get("installments", []):
            loan_sheet.append([
                row["period"],
                row["opening_balance"],
                row["principal_payment"],
                row["interest_payment"],
                row["total_payment"],
                row["closing_balance"],
            ])
        total_row = loan_sheet.max_row + 1
        first_data_row = 3
        last_data_row = total_row - 1
        loan_sheet.cell(total_row, 1, "Jami")
        for column in range(3, 6):
            letter = get_column_letter(column)
            loan_sheet.cell(total_row, column, f"=SUM({letter}{first_data_row}:{letter}{last_data_row})")
        for cell in loan_sheet[total_row]:
            cell.font = Font(name="Calibri", bold=True, color="064E3B")
            cell.fill = PatternFill("solid", fgColor="D1FAE5")
        loan_sheet.freeze_panes = "A3"
        loan_sheet.auto_filter.ref = f"A2:F{last_data_row}"

        formulas = workbook.create_sheet("Formulalar")
        _style_sheet_title(formulas, "Moliya dvigatelida qo‘llangan formulalar")
        formulas.append(["Hisob", "Formula / izoh"])
        _style_header(formulas[2])
        formula_rows = [
            ("CAPEX jami", "Uskuna xarajatlari yig‘indisi; boshlang‘ich investitsiya alohida ko‘rsatiladi."),
            ("Oylik operatsion xarajat", "Oylik ish haqi + ijara + kommunal xarajatlar."),
            ("QQS to‘lovi", "Hisoblangan chiqish QQS − hisobga olinadigan kirish QQS; natija 0 dan kichik bo‘lsa, 0."),
            ("Foyda solig‘i", "Soliqqa tortiladigan foyda × sozlangan stavka; foyda manfiy bo‘lsa, soliq bazasi 0."),
            ("Aylanmadan soliq", "Yillik tushum × sozlangan aylanma soliq stavkasi; rejim chegarasi konfiguratsiyadan olinadi."),
            ("Soddalashtirilgan QQS", "Yillik tushum × sozlangan stavka; kiruvchi QQS hisobga olinmaydi."),
            ("JShDS", "Yillik yalpi ish haqi fondi × sozlangan JShDS stavkasi."),
            ("Ijtimoiy soliq", "Yillik yalpi ish haqi fondi × sozlangan ijtimoiy soliq stavkasi."),
            ("Zararsizlik tushumi", "Oylik doimiy xarajatlar / qoplash marjasi ulushi; menyu tarkibi o‘zgarmaydi."),
            ("ROI", "Yillik sof daromad (soliqdan keyin) / boshlang‘ich investitsiya × 100."),
            ("Oylik foiz stavkasi", "Yillik foiz stavkasi / 100 / 12"),
            ("Oylik foiz to‘lovi", "Davr boshidagi qoldiq × oylik foiz stavkasi"),
            ("Annuitet to‘lovi", "P × r / (1 − (1 + r)^(-n)); r = oylik stavka"),
            ("Annuitet asosiy qarzi", "Jami oylik to‘lov − oylik foiz to‘lovi"),
            ("Differensial asosiy qarzi", "Boshlang‘ich qarz / muddat (oy)"),
            ("Jami oylik to‘lov", "Asosiy qarz to‘lovi + foiz to‘lovi"),
            ("Qoldiq oxirida", "Qoldiq boshida − asosiy qarz to‘lovi"),
            ("Hisobot qiymatlari", "Jadvaldagi to‘lovlar Python moliya dvigatelida Decimal aniqligida hisoblangan."),
        ]
        for row in formula_rows:
            formulas.append(list(row))

    _format_workbook(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _flatten(value: Any, prefix: str = "") -> list[tuple[str, Any]]:
    flattened: list[tuple[str, Any]] = []
    if isinstance(value, dict):
        for key, nested in value.items():
            label = f"{prefix}{key}"
            if isinstance(nested, (dict, list, tuple)):
                flattened.extend(_flatten(nested, f"{label} / "))
            else:
                flattened.append((label, nested))
    elif isinstance(value, (list, tuple)):
        for index, nested in enumerate(value, 1):
            flattened.extend(_flatten(nested, f"{prefix}{index} / "))
    else:
        flattened.append((prefix.rstrip(" /"), value))
    return flattened


def _excel_value(value: Any) -> Any:
    if isinstance(value, Enum):
        value = value.value
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    if value is None:
        return "Mavjud emas"
    return value


def _style_sheet_title(sheet: Any, title: str) -> None:
    sheet.append([title])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=3)
    cell = sheet.cell(1, 1)
    cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="065F46")
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26


def _style_header(cells: Any) -> None:
    for cell in cells:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="047857")
        cell.alignment = Alignment(vertical="center", wrap_text=True)


def _format_workbook(workbook: Workbook) -> None:
    thin = Side(style="thin", color="E2E8F0")
    for sheet in workbook.worksheets:
        sheet.freeze_panes = sheet.freeze_panes or "A3"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float, Decimal)):
                    cell.number_format = '#,##0.00;[Red]-#,##0.00'
        for column_cells in sheet.columns:
            column = get_column_letter(column_cells[0].column)
            max_length = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=12,
            )
            sheet.column_dimensions[column].width = min(max(max_length + 2, 14), 46)
