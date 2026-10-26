from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

from app.services.financial_exports import build_excel_report, build_pdf_report


def _sample_results() -> dict:
    return {
        "capex": {
            "equipment_items": [{"name": "=HYPERLINK(\"https://example.invalid\")", "cost": 100}],
            "equipment_total": 100,
            "initial_investment": 500,
        },
        "monthly_opex": {"payroll": 200, "rent": 100, "utilities": 50, "total": 350},
        "annual_taxes": None,
        "break_even": None,
        "roi": None,
        "loan_schedule": {
            "method": "annuity",
            "principal": 500,
            "annual_interest_rate": 12,
            "installments": [{
                "period": 1,
                "opening_balance": 500,
                "principal_payment": 40,
                "interest_payment": 5,
                "total_payment": 45,
                "closing_balance": 460,
            }],
            "total_principal": 40,
            "total_interest": 5,
            "total_payment": 45,
        },
    }


def test_pdf_export_is_a_valid_pdf_with_unicode_content() -> None:
    document = build_pdf_report(_sample_results())

    assert document.startswith(b"%PDF-")
    assert len(document) > 1_000


def test_excel_export_has_schedule_formulas_and_safe_user_strings() -> None:
    document = build_excel_report(_sample_results())
    with ZipFile(BytesIO(document)) as archive:
        assert archive.testzip() is None

    workbook = load_workbook(BytesIO(document), data_only=False)
    assert "Moliyaviy natijalar" in workbook.sheetnames
    assert "Kredit grafigi" in workbook.sheetnames
    assert "Formulalar" in workbook.sheetnames
    loan_sheet = workbook["Kredit grafigi"]
    assert loan_sheet["C4"].value == "=SUM(C3:C3)"

    summary = workbook["Moliyaviy natijalar"]
    exported_name = next(
        summary.cell(row, 3).value
        for row in range(3, summary.max_row + 1)
        if summary.cell(row, 2).value.endswith("name")
    )
    assert exported_name.startswith("'=")
